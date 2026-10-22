"""Assemble already authored semantic judgments; this script does not judge text."""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent

def read(name):
    return [json.loads(line) for line in (ROOT / name).read_text().splitlines() if line.strip()]

def write_jsonl(name, rows):
    (ROOT / name).write_text(''.join(json.dumps(row, ensure_ascii=False) + '\n' for row in rows))

cases = {r['case_no']: r for r in read('blinded_cases_256.jsonl')}
identity = {r['case_no']: r for r in read('model_identity_key.jsonl')}
manifest = json.loads((ROOT/'approved_review_manifest.json').read_text())
parts = [entry['path'] for entry in manifest['review_parts']]
for entry in manifest['review_parts']:
    assert hashlib.sha256((ROOT/entry['path']).read_bytes()).hexdigest() == entry['sha256']
assert hashlib.sha256((ROOT/'lead_overrides.jsonl').read_bytes()).hexdigest() == manifest['lead_overrides_sha256']
reviews = {}
errors = []
for name in parts:
    rows = read(name)
    assert len(rows) == 64, (name, len(rows))
    for row in rows:
        n = row['case_no']
        assert n not in reviews, n
        assert row['id'] == cases[n]['id'], (n, row['id'])
        reviews[n] = row
assert set(reviews) == set(range(1, 257))
if (ROOT / 'lead_overrides.jsonl').exists():
    for adjustment in read('lead_overrides.jsonl'):
        n = adjustment['case_no']
        assert adjustment['final_review']['id'] == cases[n]['id']
        reviews[n] = adjustment['final_review']
allowed = {'PASS', 'MINOR', 'FAIL', 'UNJUDGEABLE'}
for n, r in reviews.items():
    assert r['source_status'] in {'CLEAR', 'AMBIGUOUS', 'UNDERDETERMINED'}, n
    assert r['reference']['verdict'] in allowed, n
    for key in ['reference', 'A', 'B']:
        field = r[key]
        verdict = field['verdict'] if key == 'reference' else field['overall']
        assert verdict in allowed, (n, key, verdict)
        if key != 'reference':
            for dim in ['meaning', 'cleanup', 'fidelity']:
                assert field[dim] in allowed | {'NA'}, (n, key, dim)
            expected_fail = any(field[d] == 'FAIL' for d in ['meaning', 'cleanup', 'fidelity'])
            if expected_fail != (verdict == 'FAIL'):
                errors.append({'case_no': n, 'field': key, 'error': 'overall_fail_inconsistent'})
        if verdict == 'FAIL' and not field.get('issues'):
            errors.append({'case_no': n, 'field': key, 'error': 'failure_without_evidence'})
        text_key = 'reference' if key == 'reference' else 'output_' + key
        for item in field.get('issues', []):
            quote = item.get('quote', '')
            locations = [k for k in [text_key, 'input', 'spoken_source'] if quote and quote in cases[n][k]]
            if not locations:
                errors.append({'case_no': n, 'field': key, 'error': 'quote_not_found', 'quote': quote})
            else:
                item['quote_locations'] = locations
    assert r['winner'] in {'A', 'B', 'TIE', 'UNJUDGEABLE'}, n
    if r['winner'] in {'A','B'}:
        loser = 'B' if r['winner']=='A' else 'A'
        if r[r['winner']]['overall']=='FAIL' and r[loser]['overall'] in {'PASS','MINOR'}:
            errors.append({'case_no':n,'field':'winner','error':'failed_output_beats_usable_output'})
(ROOT / 'assembly_validation.json').write_text(json.dumps({'records': len(reviews), 'errors': errors}, ensure_ascii=False, indent=2)+'\n')
if errors:
    print(json.dumps(errors, ensure_ascii=False, indent=2))
    raise SystemExit('Judgment/evidence inconsistencies require adjudication before publishing.')

merged, curated, revised, queue = [], [], [], []
for n in range(1,257):
    c, r, key = cases[n], reviews[n], identity[n]
    m = {**c, 'source_status':r['source_status'], 'reference_review':r['reference'],
         'base_output': c['output_A' if key['A']=='base' else 'output_B'],
         'epoch2_output': c['output_A' if key['A']=='epoch2' else 'output_B'],
         'base_review': r['A' if key['A']=='base' else 'B'],
         'epoch2_review': r['A' if key['A']=='epoch2' else 'B'],
         'winner': key[r['winner']] if r['winner'] in {'A','B'} else r['winner'],
         'winner_reason': r['winner_reason'], 'identity_key': {'A':key['A'], 'B':key['B']}}
    proposal = r['reference'].get('proposed_reference')
    if r['source_status'] == 'CLEAR' and r['reference']['verdict']=='PASS':
        new_ref, origin = c['reference'], 'original_gpt_accepted'
    elif r['source_status']=='CLEAR' and isinstance(proposal,str) and proposal.strip():
        new_ref, origin = proposal, 'gpt_revised'
    else:
        new_ref, origin = None, 'quarantined'
    m['reviewed_reference'], m['reviewed_reference_origin'] = new_ref, origin
    merged.append(m)
    if new_ref is not None:
        curated.append({'case_no':n,'id':c['id'],'scene':c['scene'],'input':c['input'],
                        'reference':new_ref,'original_reference':c['reference'],
                        'reference_origin':origin,'audit_file':'audited_comparison_256.jsonl',
                        'split_role':'development_validation_previously_used_for_checkpoint_selection',
                        'human_verified':False})
        if new_ref != c['reference']:
            revised.append({'case_no':n,'id':c['id'],'input':c['input'],'old_reference':c['reference'],
                            'new_reference':new_ref,'review':r['reference']})
    else:
        queue.append(m)
clear = [r for r in merged if r['source_status']=='CLEAR']
def model_stats(rows, name):
    fields=[r[name+'_review'] for r in rows]
    semantic_judgeable=[f for f in fields if f['meaning']!='UNJUDGEABLE']
    return {'records':len(fields),'overall':dict(Counter(f['overall'] for f in fields)),
            'meaning':dict(Counter(f['meaning'] for f in fields)),
            'cleanup':dict(Counter(f['cleanup'] for f in fields)),
            'fidelity':dict(Counter(f['fidelity'] for f in fields)),
            'definite_semantic_or_identifier_failures':sum(f['meaning']=='FAIL' or f['fidelity']=='FAIL' for f in fields),
            'semantic_judgeable_records':len(semantic_judgeable),
            'overall_pass_or_minor':sum(f['overall'] in {'PASS','MINOR'} for f in fields)}
controls=[r for r in merged if r['source_status']=='CLEAR' and r['input']==r['reference'] and r['reference_review']['verdict']=='PASS']
train_audit=Path('/home/aim0/桌面/tan/ASR/Codex-web/results/refiner_training_pair_audit_200_20261022/sample_200.jsonl')
train_ids={json.loads(l)['id'] for l in train_audit.read_text().splitlines() if l.strip()}
assert not train_ids & {r['id'] for r in merged}
report = {
 'records':256,'source_status':dict(Counter(r['source_status'] for r in merged)),
 'reference_verdicts_all':dict(Counter(r['reference_review']['verdict'] for r in merged)),
 'source_clear_records':len(clear),
 'pairwise_all':dict(Counter(r['winner'] for r in merged)),
 'pairwise_source_clear':dict(Counter(r['winner'] for r in clear)),
 'models_all':{m:model_stats(merged,m) for m in ['base','epoch2']},
 'models_source_clear':{m:model_stats(clear,m) for m in ['base','epoch2']},
 'epoch2_critical_errors_by_original_reference_verdict':dict(Counter(r['reference_review']['verdict'] for r in clear if r['epoch2_review']['meaning']=='FAIL' or r['epoch2_review']['fidelity']=='FAIL')),
 'reviewed_reference_set':{'records':len(curated),'unchanged_accepted':sum(r['reference_origin']=='original_gpt_accepted' for r in curated),
                            'revised':len(revised),'quarantined':len(queue),'human_verified':False},
 'trusted_clean_controls':{'definition':'input equals original reference, source CLEAR, reference judged PASS',
                           'records':len(controls),
                           'models':{m:{'outputs_changed':sum(r[m+'_output']!=r['input'] for r in controls),
                                        'overall':dict(Counter(r[m+'_review']['overall'] for r in controls)),
                                        'definite_semantic_or_identifier_failures':sum(r[m+'_review']['meaning']=='FAIL' or r[m+'_review']['fidelity']=='FAIL' for r in controls)} for m in ['base','epoch2']}},
 'training_audit_200_id_overlap':0,
 'method':{'identity':'A/B randomized independently per case; mapping revealed only after per-case judgments',
           'review':'initial GPT review, cross review and main-model adjudication for 65–256; main review plus independent boundary review for 1–64; lead overrides recorded separately',
           'scope':'existing predictions only; no teacher generation, weight changes or training run',
           'reference':'not assumed correct; source-first assessment; digit formatting ignored but numeric values preserved'},
 'limitations':['LLM judgments and GPT-revised references are not human gold annotations.',
                'The lead reviewer had previously seen some original model comparisons; anonymized bundles reduce identity cues but do not constitute a fully double-blind study.',
                'These synthetic validation samples were previously used for checkpoint selection; this is a development audit, not an unseen real-speech test.',
                'One randomized A/B presentation per case; no full A/B order-swap consistency study was performed.',
                'Proportions describe this scene-balanced sample, not full-corpus prevalence or production performance.',
                'Source-ambiguous cases are quarantined instead of guessed; model stats on CLEAR cases use the same paired denominator.'],
 'input_hashes':{name:hashlib.sha256((ROOT/name).read_bytes()).hexdigest() for name in ['blinded_cases_256.jsonl','model_identity_key.jsonl','lead_overrides.jsonl','approved_review_manifest.json']+parts},
 'approved_review_manifest':manifest,
 'sources':{'official_evaluation':'https://github.com/AnXMuy/AgenticASR/blob/main/experiments/README.md',
            'judge_limitations':'https://arxiv.org/abs/2306.05685'}
}
for m in ['base','epoch2']:
 report['models_source_clear'][m]['issue_tags']=dict(Counter(it['tag'] for r in clear for it in r[m+'_review'].get('issues',[])))
write_jsonl('audited_comparison_256.jsonl',merged)
write_jsonl('curated_development_set.jsonl',curated)
write_jsonl('reference_revisions.jsonl',revised)
write_jsonl('manual_queue.jsonl',queue)
(ROOT/'audit_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')

columns=['case_no','id','scene','source_status','reference_verdict','input','spoken_source','reference','reviewed_reference','reference_reason',
         'base_overall','base_meaning','base_cleanup','base_fidelity','base_reason','base_output',
         'epoch2_overall','epoch2_meaning','epoch2_cleanup','epoch2_fidelity','epoch2_reason','epoch2_output','winner','winner_reason']
flat=[]
for r in merged:
 row={k:r.get(k,'') for k in columns}
 row['reference_verdict']=r['reference_review']['verdict'];row['reference_reason']=r['reference_review']['reason']
 for m in ['base','epoch2']:
  for key in ['overall','meaning','cleanup','fidelity','reason']:row[m+'_'+key]=r[m+'_review'][key]
 flat.append(row)
with (ROOT/'audit_comparison_256.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=columns);w.writeheader();w.writerows(flat)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
wb=Workbook();wb.remove(wb.active)
def sheet(name, cols, records):
 ws=wb.create_sheet(name);ws.append(cols);ws.freeze_panes='A2'
 for r in records:
  ws.append([r.get(k,'') if not isinstance(r.get(k),(dict,list)) else json.dumps(r[k],ensure_ascii=False) for k in cols])
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='1F4E78');cell.font=Font(bold=True,color='FFFFFF')
 for i,k in enumerate(cols,1):
  ws.column_dimensions[get_column_letter(i)].width=12 if k=='case_no' else (26 if k in ['id','scene','source_status','winner'] else (16 if any(t in k for t in ['verdict','overall','meaning','cleanup','fidelity']) else 80))
 for row in ws.iter_rows(min_row=2):
  ws.row_dimensions[row[0].row].height=90
  for cell in row:
   cell.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(cell.value,str):
    cell.data_type='s'
    if cell.value in ['FAIL','UNJUDGEABLE','MINOR','PASS']:
     cell.fill=PatternFill('solid',fgColor={'FAIL':'F4CCCC','UNJUDGEABLE':'FCE5CD','MINOR':'FFF2CC','PASS':'D9EAD3'}[cell.value])
 ws.auto_filter.ref=ws.dimensions
 return ws
sheet('逐条对照',columns,flat)
sheet('参考修订',['case_no','id','input','old_reference','new_reference','review'],revised)
sheet('暂缓使用',columns,[r for r in flat if r['reviewed_reference'] is None])
summary_rows=[{'item':'审计总数','value':256},{'item':'来源明确','value':len(clear)},
 {'item':'参考判定','value':report['reference_verdicts_all']},
 {'item':'参考集导出','value':report['reviewed_reference_set']},
 {'item':'明确来源的对照结果','value':report['pairwise_source_clear']}]
for m in ['base','epoch2']:summary_rows.append({'item':m,'value':report['models_source_clear'][m]})
for lim in report['limitations']:summary_rows.append({'item':'使用范围','value':lim})
sheet('汇总',['item','value'],summary_rows)
wb.save(ROOT/'audit_comparison_256.xlsx')

lines=['# 256条验证结果 GPT 语义审计','','本文件使用已保存的基座与第二轮输出；先匿名评判，完成后揭示身份。参考文本也接受独立审核。',
       '这批样本来自曾用于选择检查点的合成验证集，属于开发审计；GPT判定和修订参考不等于人工金标，也不是独立真实口语测试成绩。','',
       f'来源明确 {len(clear)} 条；其余 {256-len(clear)} 条来源存在歧义或信息不足。','',
       '## 来源明确样本上的模型比较','','| 模型 | 完全通过 | 通过或轻微问题 | 明确语义/标识错误 | 清理失败 |','|---|---:|---:|---:|---:|']
for m,label in [('base','未微调基座'),('epoch2','第二轮模型')]:
 s=report['models_source_clear'][m]
 lines.append(f"| {label} | {s['overall'].get('PASS',0)} | {s['overall_pass_or_minor']} | {s['definite_semantic_or_identifier_failures']} | {s['cleanup'].get('FAIL',0)} |")
lines+=['',f"配对胜负：{json.dumps(report['pairwise_source_clear'],ensure_ascii=False)}。",'',
        '## 参考答案处理','',f"原始判定：{json.dumps(report['reference_verdicts_all'],ensure_ascii=False)}。",
        f"导出 {len(curated)} 条GPT复核开发参考：原样保留 {len(curated)-len(revised)} 条，修订 {len(revised)} 条；隔离 {len(queue)} 条。",'',
        '修订记录保留原参考、原输入、修订文本和理由。隔离样本不加入后续开发评分集。数字格式不参与扣分，数字值与编号改变仍记录为错误。','',
        '## 第二轮模型的明确语义或标识问题样例','']
for r in [r for r in clear if r['epoch2_review']['meaning']=='FAIL' or r['epoch2_review']['fidelity']=='FAIL'][:12]:
 lines.append(f"- #{r['case_no']} `{r['id']}`：{r['epoch2_review']['reason']}")
lines+=['','## 文件','','- `audit_comparison_256.xlsx`：完整对照、参考修订及隔离表。',
        '- `audited_comparison_256.jsonl`：含全部证据的逐条审计。',
        '- `curated_development_set.jsonl`：经GPT审核或修订、来源明确的开发参考集。',
        '- `reference_revisions.jsonl`：参考修订追踪。','- `manual_queue.jsonl`：暂缓使用的歧义或无法安全修复样本。','',
        '## 依据与限制','',
        '[AgenticASR官方评估说明](https://github.com/AnXMuy/AgenticASR/blob/main/experiments/README.md)另提供基于rubric的基准评估流程，本轮属于本地开发集语义复核。',
        '[LLM裁判研究](https://arxiv.org/abs/2306.05685)指出位置、篇幅等偏差；本次匿名随机化仅降低已知身份与位置关联，不能声称消除全部裁判偏差。','']
(ROOT/'audit_summary.md').write_text('\n'.join(lines)+'\n')
print(json.dumps({k:report[k] for k in ['source_status','reference_verdicts_all','pairwise_source_clear','models_source_clear','reviewed_reference_set','trusted_clean_controls']},ensure_ascii=False,indent=2))
