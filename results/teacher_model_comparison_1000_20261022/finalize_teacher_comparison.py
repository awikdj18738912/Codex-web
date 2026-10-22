import json,hashlib,shutil,csv
from collections import Counter,defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment,PatternFill
from openpyxl.utils import get_column_letter
P=Path(__file__).resolve().parent
first=P/'teacher_generation_blind_review_200.jsonl'
first_copy=P/'teacher_generation_blind_review_200_first_pass.jsonl'
if not first_copy.exists():shutil.copyfile(first,first_copy)
rows=[json.loads(x) for x in first.read_text().splitlines() if x.strip()]
# Map blind A/B back to the two targets only after all first-pass reviews were completed.
for r in rows:
 k=r['blind_key'];w=r['blind_winner']
 r['winner']=k[w+'_origin'] if w in {'A','B'} else w
lead=[]
for r in rows:
 if r['audit_no']==122:
  old,new=r['old_review'],r['regenerated_review']
  lead.extend([
   {'audit_no':122,'id':r['id'],'field':'old_review','old':old['overall'],'new':'PASS','reason':'旧目标已经表达最终动作“把Buffer里的内容写入目标文件”；删掉说话者的显式更正过程没有改变动作事实。任务要求清理已撤回候选，不能仅因删掉纠正过程而判语义错误。','quote':'把 Buffer 里的内容写入目标文件'},
   {'audit_no':122,'id':r['id'],'field':'regenerated_review','old':new['overall'],'new':'MINOR','reason':'新目标保留最终动作，但又复述“应该写入文件，而不是服务器”的即时纠正，清理略不足，含义仍明确。','quote':'等等，应该是把 Buffer 内容写入文件，而不是写入服务器，我刚刚说反了'}])
  old.update(meaning='PASS',cleanup='PASS',fidelity='PASS',overall='PASS',issues=[],reason='最终动作写入目标文件已表达清楚；删除即时纠正过程没有改变动作事实。')
  new.update(meaning='PASS',cleanup='MINOR',fidelity='PASS',overall='MINOR',issues=[{'tag':'residual_correction','quote':'等等，应该是把 Buffer 内容写入文件，而不是写入服务器，我刚刚说反了','detail':'最终动作正确，但即时纠正过程仍有少量残留。'}],reason='最终动作正确，仍保留少量即时纠正过程。')
  r['winner']='old';r['winner_reason']='旧目标简洁表达最终动作；新目标正确但残留即时纠正。'
 if r['audit_no']==127:
  new=r['regenerated_review']
  lead.append({'audit_no':127,'id':r['id'],'field':'regenerated_review','old':new['overall'],'new':'MINOR','reason':'新目标保留订单/发票问题、客服未解决和申请动作；删去了“怕说错、要重新说一遍”的叙事背景，属于次要信息压缩，不足以判严重语义失败。旧目标更完整。','quote':'又怕说错了，所以现在还是得重新说一遍'})
  new.update(meaning='PASS',cleanup='PASS',fidelity='MINOR',overall='MINOR',issues=[{'tag':'meaningful_context_compressed','quote':'又怕说错了，所以现在还是得重新说一遍','detail':'遗漏次要叙事背景；订单/发票核心事实及申请意图保留。'}],reason='核心投诉和申请完整，轻微压缩了“怕说错、要重述”的叙事背景。')
  r['winner']='old';r['winner_reason']='旧目标保留了客服未解决后再次陈述的背景；新目标表达核心请求，但压缩掉这段次要叙事。'
(P/'lead_adjudications.jsonl').write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in lead))
merged=rows
oldc=Counter(r['old_review']['overall'] for r in merged);newc=Counter(r['regenerated_review']['overall'] for r in merged);wins=Counter(r['winner'] for r in merged)
changed=Counter('changed' if r['old_target']!=r['regenerated_target'] else 'same' for r in merged)
by=defaultdict(lambda:{'old':Counter(),'regenerated':Counter(),'winner':Counter()})
for r in merged:
 by[r['stratum']]['old'][r['old_review']['overall']]+=1;by[r['stratum']]['regenerated'][r['regenerated_review']['overall']]+=1;by[r['stratum']]['winner'][r['winner']]+=1
for r in merged:
 for field,text in [('old_review',r['old_target']),('regenerated_review',r['regenerated_target'])]:
  if r[field]['overall']=='FAIL' and not r[field]['issues']:raise ValueError(('fail_without_issue',r['audit_no'],field))
  for issue in r[field].get('issues',[]):
   if issue.get('quote') and issue['quote'] not in text and issue['quote'] not in r['input']:
    raise ValueError(('quote_not_found',r['audit_no'],field,issue['quote']))
clean=[r for r in merged if r['stratum']=='clean_or_passthrough']
gen=json.loads((P/'generation_summary.json').read_text())
report={'teacher':'Qwen3-30B-A3B-Instruct-2507-AWQ-4bit','regeneration_prompt':'current pipeline/configs/prompts_clean.json (snapshotted in generation_provenance.json)','generation':gen,
 'blind_review':{'sample_n':200,'source_status':dict(Counter(r['source_status'] for r in merged)),'risk_strata':dict(Counter(r['stratum'] for r in merged)),
  'old_target_verdicts':dict(oldc),'regenerated_target_verdicts':dict(newc),'pairwise_winners':dict(wins),'old_vs_regenerated_outputs':dict(changed),
  'verdicts_by_stratum':{s:{k:dict(v) for k,v in d.items()} for s,d in by.items()},
  'semantic_or_fidelity_failures':{'old':sum(r['old_review']['meaning']=='FAIL' or r['old_review']['fidelity']=='FAIL' for r in merged),'regenerated':sum(r['regenerated_review']['meaning']=='FAIL' or r['regenerated_review']['fidelity']=='FAIL' for r in merged)},
  'clean_controls':{'n':len(clean),'regenerated_changed':sum(r['old_target']!=r['regenerated_target'] for r in clean),'old_verdicts':dict(Counter(r['old_review']['overall'] for r in clean)),'regenerated_verdicts':dict(Counter(r['regenerated_review']['overall'] for r in clean))},'lead_adjudications':len(lead),'quote_validation_errors':0},
 'decision':'Do not replace the broader training set with same-teacher reruns yet. On this risk-stratified sample, most outputs were tied; old targets had 15 wins vs 11 for regenerated targets after main adjudication. Apply selective replacement only after a target-specific audit confirms improvement.',
 'limitations':['The sample was risk-stratified, not a prevalence estimate for the 84,875 training records.','Both targets come from the same 30B teacher; this tests regeneration with the current prompt, not a different teacher model.','The 200 judgments are model-assisted and not full human expert annotation.','The sample comes from the training split and is a data-repair diagnostic, not a student-model test set.','Regeneration used the current pipeline prompt; historical targets may have been generated under earlier prompt versions.'],
 'sample_sha256':hashlib.sha256((P/'sample_1000.jsonl').read_bytes()).hexdigest(),'prediction_sha256':hashlib.sha256((P/'predictions_qwen3_30b.jsonl').read_bytes()).hexdigest(),'first_pass_review_sha256':hashlib.sha256(first_copy.read_bytes()).hexdigest()}
(P/'teacher_generation_blind_review_200_final.jsonl').write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in merged))
(P/'teacher_generation_comparison_report_final.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
cols=['audit_no','id','scene','stratum','source_status','input','old_target','old_overall','old_meaning','old_cleanup','old_fidelity','old_reason','regenerated_target','regenerated_overall','regenerated_meaning','regenerated_cleanup','regenerated_fidelity','regenerated_reason','winner','winner_reason']
flat=[]
for r in merged:
 d={'audit_no':r['audit_no'],'id':r['id'],'scene':r['scene'],'stratum':r['stratum'],'source_status':r['source_status'],'input':r['input'],'old_target':r['old_target'],'regenerated_target':r['regenerated_target'],'winner':r['winner'],'winner_reason':r['winner_reason']}
 for prefix,name in [('old','old_review'),('regenerated','regenerated_review')]:
  for k in ['overall','meaning','cleanup','fidelity','reason']:d[prefix+'_'+k]=r[name][k]
 flat.append(d)
with (P/'teacher_generation_blind_review_200_final.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=cols);w.writeheader();w.writerows(flat)
wb=Workbook();ws=wb.active;ws.title='200条盲审';ws.append(cols);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:T{len(flat)+1}'
for d in flat:ws.append([d.get(c,'') for c in cols])
widths={'audit_no':10,'id':28,'scene':18,'stratum':26,'source_status':16,'input':80,'old_target':80,'regenerated_target':80,'old_overall':14,'old_meaning':14,'old_cleanup':14,'old_fidelity':14,'old_reason':70,'regenerated_overall':19,'regenerated_meaning':19,'regenerated_cleanup':19,'regenerated_fidelity':19,'regenerated_reason':70,'winner':18,'winner_reason':60}
for j,c in enumerate(cols,1):ws.column_dimensions[get_column_letter(j)].width=widths[c]
for row in ws.iter_rows(min_row=2):
 ws.row_dimensions[row[0].row].height=90
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
 for c in [row[7],row[12]]:
  col={'FAIL':'F4CCCC','UNJUDGEABLE':'FCE5CD','MINOR':'FFF2CC','PASS':'D9EAD3'}.get(c.value)
  if col:c.fill=PatternFill('solid',fgColor=col)
s=wb.create_sheet('汇总');s.append(['指标','结果'])
for k,v in report['blind_review'].items():s.append([k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v])
s.append(['判断',report['decision']]);s.column_dimensions['A'].width=45;s.column_dimensions['B'].width=110
for row in s.iter_rows():
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
s.freeze_panes='A2';wb.save(P/'teacher_generation_blind_review_200.xlsx')
print(json.dumps({'old':dict(oldc),'regenerated':dict(newc),'pairwise':dict(wins),'outputs':dict(changed),'clean_controls':report['blind_review']['clean_controls']},ensure_ascii=False,indent=2))
