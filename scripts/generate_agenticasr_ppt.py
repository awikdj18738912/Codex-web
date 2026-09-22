from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).resolve().parents[1] / "AgenticASR_相对原版改进与问题.pptx"
FONT = "Noto Sans CJK SC"

NAVY = RGBColor(25, 48, 82)
BLUE = RGBColor(40, 112, 205)
TEAL = RGBColor(26, 148, 139)
ORANGE = RGBColor(226, 133, 54)
RED = RGBColor(196, 75, 79)
INK = RGBColor(34, 43, 57)
MUTED = RGBColor(97, 110, 128)
LIGHT = RGBColor(247, 249, 252)
CARD = RGBColor(255, 255, 255)
LINE = RGBColor(224, 230, 238)
PALE_BLUE = RGBColor(237, 245, 255)
PALE_TEAL = RGBColor(235, 249, 247)
PALE_ORANGE = RGBColor(255, 246, 235)
PALE_RED = RGBColor(255, 241, 242)


def add_shape(slide, x, y, w, h, fill, radius=True, line=None):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_text(slide, text, x, y, w, h, size=16, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=12.5, color=INK, bullet_color=BLUE,
                gap=0.08, line_height=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for index, item in enumerate(items):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_after = Pt(gap * 12)
        p.line_spacing = line_height
        p.level = 0
        p.text = "• " + item
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
    return box


def add_header(slide, kicker, title, subtitle, page):
    add_text(slide, kicker.upper(), 0.62, 0.34, 4.2, 0.24, size=9.5, color=BLUE, bold=True)
    add_text(slide, title, 0.62, 0.66, 11.6, 0.55, size=25, color=NAVY, bold=True)
    add_text(slide, subtitle, 0.64, 1.28, 11.2, 0.32, size=11.5, color=MUTED)
    add_shape(slide, 0.64, 1.70, 12.05, 0.025, BLUE, radius=False)
    add_text(slide, f"0{page}", 12.15, 0.42, 0.5, 0.25, size=10, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)


def add_card_title(slide, title, x, y, w, accent=BLUE, subtitle=None):
    add_shape(slide, x, y, w, 0.06, accent, radius=False)
    add_text(slide, title, x + 0.18, y + 0.18, w - 0.36, 0.30, size=15, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, subtitle, x + 0.18, y + 0.54, w - 0.36, 0.40, size=10.5, color=MUTED)


def add_tag(slide, text, x, y, w, fill=PALE_BLUE, color=BLUE):
    add_shape(slide, x, y, w, 0.31, fill, radius=True)
    add_text(slide, text, x, y + 0.01, w, 0.26, size=9.5, color=color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.01)


def slide_one(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(
        slide,
        "AgenticASR · 项目进展",
        "相对于原始 AgenticASR 的系统改进",
        "从“流式精修原型”扩展为“可控、可审计的 GPU 服务系统”",
        1,
    )

    # Baseline column.
    add_shape(slide, 0.64, 2.05, 3.25, 4.55, CARD, line=LINE)
    add_card_title(slide, "原始 AgenticASR 基线", 0.64, 2.05, 3.25, accent=MUTED)
    add_tag(slide, "VAD → ASR → K=3 Refiner", 0.84, 2.78, 2.45, fill=RGBColor(241, 243, 247), color=MUTED)
    add_bullets(
        slide,
        [
            "主要链路：VAD、sherpa-onnx 在线 ASR、ChunkManager、滑动窗口 Refiner。",
            "重点解决口语清理、自我修正和书面化表达。",
            "原版已经关注术语与实体，但主要通过提示词和训练数据表达。",
        ],
        0.86, 3.30, 2.82, 1.95, size=12.2, bullet_color=MUTED,
    )
    add_shape(slide, 0.84, 5.54, 2.45, 0.70, RGBColor(249, 250, 252), line=LINE)
    add_text(slide, "当前改造的重点：\n把模型能力下沉为可验证的运行时管线。",
             1.00, 5.68, 2.12, 0.40, size=10.5, color=NAVY, bold=True)

    # Improvements column.
    add_shape(slide, 4.12, 2.05, 8.56, 4.55, CARD, line=LINE)
    add_card_title(slide, "当前版本新增能力", 4.12, 2.05, 8.56, accent=BLUE,
                   subtitle="四个方向：实体保护、服务扩展、精修控制、安全校验")

    cards = [
        (4.36, 2.90, 3.84, 1.35, PALE_BLUE, BLUE, "实体 / 术语保护",
         ["SQLite 实体库、标准名、别名和领域管理", "精确/受控模糊匹配 → 占位保护 → 恢复校验", "降低人名、地名、作品名被改写的风险"]),
        (8.36, 2.90, 4.08, 1.35, PALE_TEAL, TEAL, "ASR 与服务架构",
         ["扩展 Qwen3-ASR GPU + Transformer Refiner", "支持 Web、文件、离线流式和实时流式", "ASR 与 Refiner 可独立持续运行"]),
        (4.36, 4.43, 3.84, 1.35, PALE_ORANGE, ORANGE, "精修调度控制",
         ["增加 off / conservative / tri_state 门控", "KEEP / DEFER / REFINE 三状态", "窗口聚合后统一复核，减少跨窗口错断"]),
        (8.36, 4.43, 4.08, 1.35, RGBColor(242, 244, 250), NAVY, "安全与可审计",
         ["数字、量词、成语、标点和内容完整性校验", "结构化 JSON 局部修改", "失败重试、回退和 JSONL 审计日志"]),
    ]
    for x, y, w, h, fill, accent, title, bullets in cards:
        add_shape(slide, x, y, w, h, fill, line=fill)
        add_text(slide, title, x + 0.16, y + 0.13, w - 0.32, 0.24, size=12.2, color=accent, bold=True)
        add_bullets(slide, bullets, x + 0.15, y + 0.42, w - 0.30, h - 0.48,
                    size=9.2, color=INK, bullet_color=accent, gap=0.01, line_height=0.91)

    add_shape(slide, 0.64, 6.82, 12.04, 0.42, NAVY)
    add_text(slide, "核心变化：从“模型直接改文本”转向“模型精修 + 实体保护 + 多层校验”。",
             0.88, 6.90, 11.55, 0.25, size=12.2, color=RGBColor(255, 255, 255), bold=True,
             valign=MSO_ANCHOR.MIDDLE)


def slide_two(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    add_header(
        slide,
        "AgenticASR · 问题与研究",
        "改造后的问题与原因",
        "功能覆盖提升，但模块耦合、判断不稳定和处理成本也随之增加",
        2,
    )

    # Problems.
    add_shape(slide, 0.64, 2.05, 5.75, 4.55, CARD, line=LINE)
    add_card_title(slide, "当前暴露的问题", 0.64, 2.05, 5.75, accent=RED,
                   subtitle="当前回归与日志中反复出现的问题")
    problem_items = [
        ("叠字漏修", "正常叠词与口吃叠字难以区分，部分重复未清除，模型判断也可能不一致。"),
        ("断句错位", "句号、逗号位置仍可能错误，导致句子被拆开或多个句子发生粘连。"),
        ("保护与修正冲突", "实体、数字、成语保护和 tri_state 门控可能阻止模型的正确修改。"),
        ("成本增加", "局部复核带来额外 GPU 调用；模型格式异常时会触发回退。"),
    ]
    y = 2.88
    for title, body in problem_items:
        add_text(slide, title, 0.92, y, 1.25, 0.25, size=11.5, color=RED, bold=True)
        add_text(slide, body, 2.03, y, 3.98, 0.48, size=10.7, color=INK)
        y += 0.78

    # Causes.
    add_shape(slide, 6.64, 2.05, 6.04, 4.55, CARD, line=LINE)
    add_card_title(slide, "问题成因", 6.64, 2.05, 6.04, accent=TEAL,
                   subtitle="当前问题主要来自模型歧义与多模块协同")
    cause_items = [
        ("01", "上下文不足", "流式窗口只提供局部上下文，叠字和句子边界缺少完整语义。"),
        ("02", "模块耦合", "VAD、ASR 分段、窗口聚合和标点恢复会相互影响。"),
        ("03", "语义有歧义", "同样的重复形式可能是正常词语，也可能是口吃或重复起句。"),
        ("04", "校验偏保守", "保护和安全校验拒绝高风险修改时，会同时造成部分漏修。"),
        ("05", "协议不稳定", "模型偶尔不按结构化格式返回，导致重试或回退。"),
    ]
    y = 2.82
    for num, title, body in cause_items:
        add_shape(slide, 6.94, y + 0.02, 0.42, 0.28, PALE_TEAL)
        add_text(slide, num, 6.94, y + 0.02, 0.42, 0.25, size=9.3, color=TEAL, bold=True,
                 align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.01)
        add_text(slide, title, 7.50, y, 1.30, 0.25, size=11.3, color=NAVY, bold=True)
        add_text(slide, body, 8.70, y, 3.55, 0.42, size=10.2, color=INK)
        y += 0.66

    add_shape(slide, 0.64, 6.82, 12.04, 0.42, TEAL)
    add_text(slide, "当前结论：系统的安全性和可控性提升了，但精修准确率仍受到上下文、规则和模型稳定性的共同影响。",
             0.88, 6.90, 11.55, 0.25, size=12.0, color=RGBColor(255, 255, 255), bold=True,
             valign=MSO_ANCHOR.MIDDLE)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_one(prs)
    slide_two(prs)
    prs.core_properties.title = "AgenticASR 相对原版的改进与问题"
    prs.core_properties.subject = "两页项目进展汇报"
    prs.core_properties.author = "Codex"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
